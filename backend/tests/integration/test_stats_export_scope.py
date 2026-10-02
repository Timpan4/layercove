"""Public CSV/XLSX statistics exports retain selected dates and filters."""

import csv
import io
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app.models.project import Project
from backend.app.models.user import User


def export_rows(response, format):
    if format == "csv":
        return list(csv.reader(io.StringIO(response.text)))
    workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    return [list(row) for row in workbook.active.iter_rows(values_only=True)]


@pytest.fixture
async def export_scope_data(db_session, printer_factory, archive_factory):
    alice = User(username="export-alice")
    bob = User(username="export-bob")
    project = Project(name="Export fixture")
    other_project = Project(name="Other fixture")
    db_session.add_all([alice, bob, project, other_project])
    await db_session.commit()
    printer = await printer_factory()
    other_printer = await printer_factory()
    now = datetime.now(timezone.utc)
    records = [
        (datetime(2020, 1, 5, tzinfo=timezone.utc), "completed", alice.id, printer.id, project.id),
        (datetime(2020, 1, 12, 23, 59, 59, tzinfo=timezone.utc), "failed", alice.id, printer.id, project.id),
        (datetime(2020, 1, 4, 23, 59, 59, tzinfo=timezone.utc), "failed", alice.id, printer.id, project.id),
        (datetime(2020, 1, 13, tzinfo=timezone.utc), "completed", alice.id, printer.id, project.id),
        (datetime(2020, 1, 8, tzinfo=timezone.utc), "failed", bob.id, printer.id, project.id),
        (datetime(2020, 1, 9, tzinfo=timezone.utc), "completed", alice.id, other_printer.id, project.id),
        (datetime(2020, 1, 10, tzinfo=timezone.utc), "completed", alice.id, printer.id, other_project.id),
        (datetime(2020, 1, 11, tzinfo=timezone.utc), "failed", None, printer.id, project.id),
        (now - timedelta(days=10), "completed", alice.id, printer.id, project.id),
        (now - timedelta(days=60), "failed", alice.id, printer.id, project.id),
    ]
    for created_at, status, user_id, printer_id, project_id in records:
        await archive_factory(
            printer_id=printer_id,
            project_id=project_id,
            created_by_id=user_id,
            created_at=created_at,
            status=status,
            failure_reason="Fixture failure" if status == "failed" else None,
        )
    return {"user": alice.id, "printer": printer.id, "project": project.id, "today": now.date().isoformat()}


@pytest.mark.asyncio
@pytest.mark.integration
@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize("scope", ["today", "historical", "all-time", "no-user", "legacy-default", "legacy-days"])
async def test_public_statistics_export_scope(async_client, export_scope_data, format, scope):
    data = export_scope_data
    params = {"format": format}
    total, failed = 0, 0
    if scope == "today":
        params.update(date_from=data["today"], date_to=data["today"])
    elif scope in ("historical", "no-user"):
        params.update(
            date_from="2020-01-05",
            date_to="2020-01-12",
            created_by_id=data["user"] if scope == "historical" else -1,
            printer_id=data["printer"],
            project_id=data["project"],
        )
        total, failed = (2, 1) if scope == "historical" else (1, 1)
    elif scope == "all-time":
        params["all_time"] = "true"
        total, failed = 10, 5
    elif scope == "legacy-default":
        total, failed = 1, 0
    else:
        params["days"] = 90
        total, failed = 2, 1

    response = await async_client.get("/api/v1/archives/stats/export", params=params)
    assert response.status_code == 200
    assert f'.{format}"' in response.headers["content-disposition"]
    artifact_dir = os.environ.get("STATS_EXPORT_ARTIFACT_DIR")
    if artifact_dir:
        Path(artifact_dir).mkdir(parents=True, exist_ok=True)
        (Path(artifact_dir) / f"{scope}.{format}").write_bytes(response.content)
    rows = export_rows(response, format)
    metrics = {row[0]: row[1] for row in rows if len(row) > 1 and row[0]}
    assert int(metrics["Total Prints"]) == total
    assert int(metrics["Failed Prints"]) == failed

    if scope == "all-time":
        assert metrics["Period"] == "All time"
        assert "Period (days)" not in metrics
    elif scope in ("today", "historical", "no-user"):
        assert metrics["Start date"] == params["date_from"]
        assert metrics["End date"] == params["date_to"]
        trend_header = next(i for i, row in enumerate(rows) if row[0] == "Week")
        trend = [row for row in rows[trend_header + 1 :] if row[0]]
        assert sum(int(row[1] or 0) for row in trend) == total
        assert sum(int(row[2] or 0) for row in trend) == failed
        assert all(params["date_from"] <= row[0] <= params["date_to"] for row in trend)
    else:
        assert int(metrics["Period (days)"]) == (30 if scope == "legacy-default" else 90)
